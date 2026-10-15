"""Narrow editorial and table-fit corrections after the v13 contents refresh."""
from pathlib import Path
import argparse
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
p=argparse.ArgumentParser();p.add_argument('document',type=Path);a=p.parse_args()
d=Document(a.document)
for para in d.paragraphs:
    if para.text.startswith('22.1 Lessons from implementation\t'):
        for text in para._p.iter(qn('w:t')):
            if text.text:
                text.text=text.text.replace('Lessons from implementation','Lessons')
    for run in para.runs:
        old=run.text
        new=old.replace('Session and audit metadata confirm closure, released capacity and automation cleanup.', 'Earlier release tests confirmed session closure, released capacity and automation cleanup.')
        new=new.replace('Authenticated browser checks passed session start,', 'Earlier release browser checks passed session start,')
        if new!=old:run.text=new
for table in d.tables:
    for row in table.rows:
        for cell in row.cells:
            if cell.text=='Escalation':
                for para in cell.paragraphs:
                    for run in para.runs:run.font.size=Pt(9)
d.save(a.document)
